"""Turning a user's spreadsheets into tables the agent can query.

The agent itself only ever holds a read-only connection. This module is the one
place that writes, and it is ordinary application code - the SQL here is built
by us from sanitised names, never by the model - so the guardrails that police
the agent's SQL do not apply to it.
"""

from __future__ import annotations

import csv
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path

import duckdb

SUPPORTED_SUFFIXES = {".csv", ".tsv", ".xlsx"}

# Names the agent is likely to confuse with SQL, or that DuckDB reserves.
_RESERVED = {
    "select", "from", "where", "table", "order", "group", "by", "join", "limit",
    "user", "index", "key", "values", "date", "time", "data",
}


class DatasetError(ValueError):
    """A problem with a file the user can do something about. The message is
    shown to them as-is, so it says what to do."""


@dataclass(slots=True)
class ImportResult:
    table: str
    rows: int
    columns: int


def table_name_for(filename: str) -> str:
    """A safe SQL identifier from a filename: `Sales Q3 (final).xlsx` ->
    `sales_q3_final`. Also what makes re-uploading a file replace its table."""
    stem = Path(filename).stem.lower()
    name = re.sub(r"[^a-z0-9]+", "_", stem).strip("_") or "table"
    if name[0].isdigit():
        name = f"t_{name}"
    if name in _RESERVED:
        name = f"{name}_table"
    return name[:60]


def _xlsx_to_csv(path: Path, out: Path) -> None:
    """First sheet only. openpyxl rather than pandas: a few hundred KB instead
    of tens of MB, for a job that is reading cells in order."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - it is in requirements
        raise DatasetError("Reading Excel files needs `pip install openpyxl`.") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise DatasetError("the file could not be opened as an Excel workbook") from exc
    sheet = workbook.worksheets[0]
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() for cell in row):
                writer.writerow(["" if cell is None else cell for cell in row])
    workbook.close()


def import_file(path: Path, db_path: Path) -> ImportResult:
    """Load one CSV/TSV/XLSX file into `db_path` as a table named after it.

    DuckDB infers column types, which is what lets the agent do arithmetic on
    a column that arrived as text in a spreadsheet.
    """
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DatasetError(f"unsupported type - use {', '.join(sorted(SUPPORTED_SUFFIXES))}")

    table = table_name_for(path.name)
    db_path.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmp:
        source = path
        if suffix == ".xlsx":
            source = Path(tmp) / "sheet.csv"
            _xlsx_to_csv(path, source)

        connection = duckdb.connect(str(db_path))
        try:
            delimiter = "\t" if suffix == ".tsv" else None
            reader = (
                "read_csv_auto(?, header = true, delim = ?)"
                if delimiter
                else "read_csv_auto(?, header = true)"
            )
            params = [str(source), delimiter] if delimiter else [str(source)]
            try:
                connection.execute(
                    f'CREATE OR REPLACE TABLE "{table}" AS SELECT * FROM {reader}', params
                )
            except duckdb.Error as exc:
                raise DatasetError(
                    "the file could not be read as a table - check it has a header row "
                    "and one record per line"
                ) from exc
            rows = connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
            columns = len(connection.execute(f'DESCRIBE "{table}"').fetchall())
            if rows == 0:
                connection.execute(f'DROP TABLE "{table}"')
                raise DatasetError("the file has a header but no rows")
            connection.execute("CHECKPOINT")
        finally:
            connection.close()
    return ImportResult(table=table, rows=int(rows), columns=int(columns))


def drop_table(table: str, db_path: Path) -> bool:
    if not db_path.exists():
        return False
    connection = duckdb.connect(str(db_path))
    try:
        existing = {
            row[0] for row in connection.execute(
                "SELECT table_name FROM information_schema.tables WHERE table_schema = 'main'"
            ).fetchall()
        }
        if table not in existing:
            return False
        connection.execute(f'DROP TABLE "{table}"')
        connection.execute("CHECKPOINT")
        return True
    finally:
        connection.close()


def table_count(db_path: Path) -> int:
    if not db_path.exists():
        return 0
    connection = duckdb.connect(str(db_path), read_only=True)
    try:
        return connection.execute(
            "SELECT COUNT(*) FROM information_schema.tables WHERE table_schema = 'main'"
        ).fetchone()[0]
    finally:
        connection.close()
